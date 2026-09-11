"""
Excel (.xlsx) export helpers for Modbus polling logs and IEC 104 event logs.
Uses openpyxl. Produces a styled, readable workbook with a header row,
frozen panes, autosized columns, and light banding.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F3B57", end_color="1F3B57", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BAND_FILL = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _style_sheet(ws, headers, rows):
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for r, row in enumerate(rows, start=2):
        ws.append(row)
        band = (r % 2 == 0)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN_BORDER
            if band:
                cell.fill = BAND_FILL

    # autosize columns (rough heuristic based on content length)
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_len = max(max_len, len(str(row[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 10), 60)

    ws.freeze_panes = "A2"


def export_modbus_log_to_excel(path, log_rows, meta=None):
    """
    log_rows: list of (timestamp, start_address, [values]) as collected by the poller.
    meta: optional dict of connection info to write to a summary sheet.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Modbus Log"

    max_vals = max((len(v) for _, _, v in log_rows), default=0)
    headers = ["Timestamp", "Start Address"] + [f"Reg[{i}]" for i in range(max_vals)]
    rows = []
    for ts, addr, values in log_rows:
        row = [ts, addr] + list(values) + [""] * (max_vals - len(values))
        rows.append(row)

    _style_sheet(ws, headers, rows)

    if meta:
        ws2 = wb.create_sheet("Connection Info")
        ws2.append(["Parameter", "Value"])
        for c in (1, 2):
            cell = ws2.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for k, v in meta.items():
            ws2.append([k, v])
        ws2.column_dimensions["A"].width = 24
        ws2.column_dimensions["B"].width = 40

    wb.save(path)


def export_table_snapshot_to_excel(path, headers, rows, sheet_name="Values"):
    """Export whatever is currently rendered in the values table (a single snapshot)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    _style_sheet(ws, headers, rows)
    wb.save(path)


def export_iec104_log_to_excel(path, log_rows, meta=None):
    """
    log_rows: list of dicts with keys like timestamp, common_address, io_address,
    type, value, quality, cot (cause of transmission).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "IEC 104 Log"

    headers = ["Timestamp", "Common Address", "IO Address", "Type", "Value", "Quality", "COT"]
    rows = []
    for entry in log_rows:
        rows.append([
            entry.get("timestamp", ""),
            entry.get("common_address", ""),
            entry.get("io_address", ""),
            entry.get("type", ""),
            entry.get("value", ""),
            entry.get("quality", ""),
            entry.get("cot", ""),
        ])

    _style_sheet(ws, headers, rows)

    if meta:
        ws2 = wb.create_sheet("Connection Info")
        ws2.append(["Parameter", "Value"])
        for c in (1, 2):
            cell = ws2.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for k, v in meta.items():
            ws2.append([k, v])
        ws2.column_dimensions["A"].width = 24
        ws2.column_dimensions["B"].width = 40

    wb.save(path)
